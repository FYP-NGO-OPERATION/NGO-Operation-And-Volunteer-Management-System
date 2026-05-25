import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

prs = Presentation()

# Layouts: 0 is Title Slide, 1 is Title and Content, 5 is Title Only, 6 is Blank

diagrams_dir = r"E:\Fyp\ngo_volunteer_app\docs\diagrams"

def add_slide(title_text, body_text="", img_file=None):
    if body_text and not img_file:
        slide = prs.slides.add_slide(prs.slide_layouts[1]) # Title and Content
        title = slide.shapes.title
        title.text = title_text
        body = slide.shapes.placeholders[1]
        tf = body.text_frame
        tf.text = body_text
    elif img_file:
        slide = prs.slides.add_slide(prs.slide_layouts[5]) # Title Only
        title = slide.shapes.title
        title.text = title_text
        img_path = os.path.join(diagrams_dir, img_file)
        if os.path.exists(img_path):
            # Center image
            left = Inches(1)
            top = Inches(1.5)
            width = Inches(8)
            # Try to add picture, if it's too tall we might need height constraints but default is ok
            slide.shapes.add_picture(img_path, left, top, width=width)
        
        # If there's body text with the image, add a textbox
        if body_text:
            txBox = slide.shapes.add_textbox(Inches(0.5), Inches(6), Inches(9), Inches(1.5))
            tf = txBox.text_frame
            tf.text = body_text
            tf.word_wrap = True
    else:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title = slide.shapes.title
        title.text = title_text
    
    return slide

# Slide 1: Title Page
slide = prs.slides.add_slide(prs.slide_layouts[0])
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "AI-Integrated Operations & Volunteer Coordination Platform"
subtitle.text = "Final Year Project - I Proposal Defense\nSubmitted by: Muhammad Maauz Mansoor (233599)\nSupervisor: Mam Fatima Yousuf"

# 1. Project Title
add_slide("1. Project Title", "AI-Integrated Operations and Volunteer Coordination Platform for Charitable NGOs")

# 2. Project Overview
overview_text = ("• NGOs (like HRAS) are reliant on scattered WhatsApp groups and handwritten registers.\n"
                 "• This manual approach causes operational friction, lost financial data, and unorganized volunteer tracking.\n"
                 "• Proposed Solution: A comprehensive, AI-integrated platform built with Flutter acting as a centralized hub to handle campaigns, track volunteers, and log donations transparently.")
add_slide("2. Project Overview", overview_text)

# 3. Project Goals and Objectives
goals_text = ("• Create a unified digital ecosystem that seamlessly handles the NGO's operations.\n"
              "• Establish complete operational transparency by centralizing real-time tracking for donations and expenses.\n"
              "• Introduce accessible AI capabilities for actionable insights and automated reporting.\n"
              "• Deliver a cross-platform experience (Web, Android, iOS) from a single codebase.")
add_slide("3. Project Goals and Objectives", goals_text)

# 4. High-Level System Components
components_text = ("1. Cross-Platform Mobile Application (Flutter)\n"
                   "2. Administrative Web Dashboard (Flutter Web)\n"
                   "3. Real-Time Cloud Infrastructure (Firebase)\n"
                   "4. Intelligent Reporting Engine (AI)")
add_slide("4. High-Level System Components", components_text)

# 4b. High-Level System Components (Concept Map Diagram)
add_slide("4. High-Level System Components (Concept Map)", "", "ch1_workflow.png")

# 5. Optional Functional Units
opt_text = ("• Cryptographic QR Attendance: For quick volunteer verification.\n"
            "• Dynamic PDF Receipt Generation: Automated receipts for donors.\n"
            "• Smart Matching Algorithm: Suggesting campaigns based on volunteer skills.")
add_slide("5. Optional Functional Units", opt_text)

# 6. Exclusions
excl_text = ("• Live Payment Gateway Integration: Donations will be logged manually or via uploaded receipts to avoid complex corporate banking compliance during MVP.\n"
             "• Automated Background Checks: Manual verification will be used for volunteers.\n"
             "• Direct Social Media APIs: Automated posting to Facebook/Twitter is excluded.")
add_slide("6. Exclusions", excl_text)

# 7. Application Architecture
arch_text = ("The system utilizes a 4-Tier Architecture:\n"
             "• Presentation Layer: Flutter App & Web Dashboard\n"
             "• AI Services Layer: Firebase ML Kit\n"
             "• Backend Services: Firebase Auth & Cloud Functions\n"
             "• Data Layer: Cloud Firestore (NoSQL)")
add_slide("7. Application Architecture", arch_text)

# 7b. Application Architecture Diagram
add_slide("7. Application Architecture (Diagram)", "", "ch5_component_arch.png")

# 8. Gantt Chart
add_slide("8. Gantt Chart (Full FYP Timeline)", "", "full_fyp_timeline.png")

# 9. Hardware and Software Specification
hw_sw_text = ("Hardware Requirements:\n"
              "• Standard PC/Mac for development (8GB RAM minimum).\n"
              "• Android/iOS device for physical testing.\n\n"
              "Software Requirements:\n"
              "• Flutter SDK & Dart\n"
              "• Android Studio / VS Code\n"
              "• Firebase CLI")
add_slide("9. Hardware and Software Specification", hw_sw_text)

# 10. Tools and Technologies (With Reasoning)
tools_text = ("• Flutter & Dart: Chosen for single-codebase cross-platform deployment, saving immense development time.\n"
              "• Google Firebase (Firestore & Auth): Chosen as the Backend-as-a-Service (BaaS) for superior real-time data synchronization.\n"
              "• Firebase ML Kit & Python: Utilized for backend AI analytics and reporting.")
add_slide("10. Tools and Technologies (With Reasoning)", tools_text)

# 11. Expertise of Team Members
team_text = ("Team Member: Muhammad Maauz Mansoor (233599)\n\n"
             "Relevant Expertise:\n"
             "• Flutter UI Development\n"
             "• Database Schema Design (NoSQL)\n"
             "• Object-Oriented Programming (OOP)\n"
             "• AI Integration Strategies")
add_slide("11. Expertise of Team Members", team_text)

# 12. References
ref_text = ("[1] S. Newman, 'Building Microservices', O'Reilly Media, 2015.\n"
            "[2] Google, 'Flutter Architectural Overview', [Online].\n"
            "[3] Firebase Documentation, 'Understand Cloud Firestore', [Online].")
add_slide("12. References", ref_text)

prs.save(r'E:\Fyp\FYP_Proposal_Presentation_Final.pptx')
print("Successfully generated final presentation.")
