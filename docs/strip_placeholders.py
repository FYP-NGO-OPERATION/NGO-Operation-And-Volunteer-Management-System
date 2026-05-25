import os
import glob
from docx import Document

def strip_placeholders(docx_path):
    try:
        doc = Document(docx_path)
        modified = False
        
        # We need to remove paragraphs that contain the literal text [PLACEHOLDER
        # A simple approach is to clear the text of those paragraphs.
        for p in doc.paragraphs:
            if '[PLACEHOLDER' in p.text:
                p.clear()
                modified = True
                
        if modified:
            doc.save(docx_path)
            print(f"Cleaned placeholders from: {os.path.basename(docx_path)}")
        else:
            print(f"No placeholders found in: {os.path.basename(docx_path)}")
            
    except Exception as e:
        print(f"Error processing {os.path.basename(docx_path)}: {e}")

def main():
    base_path = r"E:\Fyp\ngo_volunteer_app\docs"
    
    # Get all docx files recursively
    docx_files = glob.glob(os.path.join(base_path, "**", "*.docx"), recursive=True)
    
    if not docx_files:
        print("No docx files found.")
        return
        
    for f in docx_files:
        # Ignore temp files
        if os.path.basename(f).startswith('~$'):
            continue
        strip_placeholders(f)
        
if __name__ == "__main__":
    main()
